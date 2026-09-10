"""Prepare privacy-minimized website data and independent SciPy reference fixtures.

Usage: python prepare_survey.py --input path/to/survey.xlsx --output path/to/data
The source workbook is read-only; output records retain ineligible ratings so the
question-specific eligibility rule remains explicit at every analysis boundary.
"""

import argparse
from collections import Counter
import hashlib
import importlib.metadata
import json
from pathlib import Path
import platform
import re
import unicodedata

import numpy as np
import openpyxl
from scipy import stats


VERSION = "1.0.0"
SEED = 20260909
PROVIDERS = ("Airtel", "Jio")
PLACEHOLDER = "No service failure is included in this persona."
METRIC_SPECS = [
    ("reliability", "Overall reliability", "I", "all", "Reliability"),
    ("coverage", "Network coverage", "J", "all", "Reliability"),
    ("speed", "Speed consistency", "K", "all", "Reliability"),
    ("promise", "Promise delivery", "L", "all", "Reliability"),
    ("access", "Support accessibility", "N", "support", "Responsiveness"),
    ("response", "Response speed", "O", "support", "Responsiveness"),
    ("resolution", "Final resolution", "P", "support", "Responsiveness"),
    ("assurance", "Representative competence", "Q", "support", "Assurance"),
    ("empathy", "Understanding the problem", "R", "support", "Empathy"),
    ("digital", "Digital self-service", "S", "all", "Tangibles (digital proxy)"),
    ("recovery", "Recovery effectiveness", "W", "failure", None),
    ("fairness", "Complaint fairness", "Y", "failure", None),
    ("continue", "Likelihood to continue", "AC", "all", None),
    ("recommend", "Likelihood to recommend", "AD", "all", None),
]
HEADLINE = ("reliability", "response", "assurance", "empathy", "digital")
POPULATIONS = {metric[0]: metric[3] for metric in METRIC_SPECS}


def normalize(value):
    return " ".join(unicodedata.normalize("NFKC", str(value or "")).split())


def selections(value):
    return list(dict.fromkeys(part for token in normalize(value).split(";") if (part := normalize(token))))


def boolean(value, field, row):
    cleaned = normalize(value)
    if cleaned not in ("Yes", "No"):
        raise ValueError(f"Row {row}, {field}: expected Yes or No, got {cleaned!r}")
    return cleaned == "Yes"


def rating(value, column, row):
    if value is None or normalize(value) == "":
        return None
    if isinstance(value, bool) or not isinstance(value, (int, float)) or value not in (1, 2, 3, 4, 5):
        raise ValueError(f"Row {row}, {column}: rating outside integer scale 1–5: {value!r}")
    return int(value)


def eligible(record, metric):
    population = POPULATIONS[metric]
    return population == "all" or record[population]


def values(records, metric, provider=None):
    return [r["ratings"][metric] for r in records if (provider is None or r["provider"] == provider)
            and eligible(r, metric) and r["ratings"][metric] is not None]


def summary(sequence):
    sequence = np.asarray(sequence, dtype=float)
    n = len(sequence)
    counts = [int(np.count_nonzero(sequence == score)) for score in range(1, 6)]
    return {"n": n, "mean": float(np.mean(sequence)) if n else None,
            "median": float(np.median(sequence)) if n else None,
            "modes": [i + 1 for i, count in enumerate(counts) if count == max(counts)] if n else [],
            "variance": float(np.var(sequence, ddof=1)) if n > 1 else None,
            "sd": float(np.std(sequence, ddof=1)) if n > 1 else None,
            "counts": counts, "proportions": [c / n for c in counts] if n else [None] * 5}


def comparison(a, b):
    result = {"Airtel": summary(a), "Jio": summary(b), "u": None, "p": None,
              "cliffsDelta": None, "meanDifference": None, "inferenceAvailable": min(len(a), len(b)) >= 20}
    if not a or not b:
        return result
    if len(set(a + b)) == 1:
        result.update(u=len(a) * len(b) / 2, cliffsDelta=0.0, meanDifference=0.0,
                      inferenceAvailable=False, unavailableReason="All eligible ratings are identical; rank variance is zero.")
        return result
    test = stats.mannwhitneyu(a, b, alternative="two-sided", method="asymptotic", use_continuity=True)
    result.update(u=float(test.statistic), p=float(test.pvalue),
                  cliffsDelta=float(2 * test.statistic / (len(a) * len(b)) - 1),
                  meanDifference=float(np.mean(a) - np.mean(b)))
    return result


def cronbach(records, metrics):
    rows = [[r["ratings"][key] for key in metrics] for r in records
            if all(eligible(r, key) and r["ratings"][key] is not None for key in metrics)]
    matrix = np.asarray(rows, dtype=float)
    if len(rows) < 2:
        return {"n": len(rows), "items": metrics, "alpha": None}
    k = len(metrics)
    total_variance = np.var(matrix.sum(axis=1), ddof=1)
    alpha = k / (k - 1) * (1 - np.var(matrix, axis=0, ddof=1).sum() / total_variance) if total_variance else None
    return {"n": len(rows), "items": metrics, "alpha": float(alpha) if alpha is not None else None}


def matches(record, filters):
    for key, selected in filters.items():
        if key == "services":
            if selected and not any(service in record["services"] for service in selected):
                return False
        elif isinstance(selected, list):
            if selected and record[key] not in selected:
                return False
        elif record[key] != selected:
            return False
    return True


def holm(results):
    ordered = sorted(((key, result["p"]) for key, result in results.items()
                      if key in HEADLINE and result["inferenceAvailable"]), key=lambda pair: pair[1])
    prior = 0
    adjusted = {}
    for index, (key, p) in enumerate(ordered):
        # The declared family remains five, including any unavailable proxy.
        prior = max(prior, min(1.0, (5 - index) * p))
        adjusted[key] = prior
    return adjusted


def categorical(records, field):
    table = [[sum(r["provider"] == provider and r[field] == answer for r in records)
              for answer in (False, True)] for provider in PROVIDERS]
    result = {"field": field, "rows": list(PROVIDERS), "columns": ["No", "Yes"], "table": table,
              "n": len(records), "chi2": None, "p": None, "cramersV": None, "expected": None,
              "inferenceAvailable": False}
    if not all(sum(row) for row in table) or not all(sum(row[i] for row in table) for i in range(2)):
        return result
    chi2, p, dof, expected = stats.chi2_contingency(table, correction=False)
    result.update(chi2=float(chi2), p=float(p), df=int(dof), expected=expected.tolist(),
                  cramersV=float(np.sqrt(chi2 / len(records))),
                  inferenceAvailable=bool(np.min(expected) >= 5 and min(map(sum, table)) >= 20))
    return result


def relationship(records, x_key, y_key):
    pairs = [(r["ratings"][x_key], r["ratings"][y_key]) for r in records
             if eligible(r, x_key) and eligible(r, y_key)
             and r["ratings"][x_key] is not None and r["ratings"][y_key] is not None]
    x, y = zip(*pairs)
    result = stats.spearmanr(x, y)
    return {"x": x_key, "y": y_key, "n": len(pairs), "rho": float(result.statistic),
            "asymptoticP": float(result.pvalue),
            "note": "asymptoticP is a reference diagnostic; Studio uses 4,999-permutation inference."}


def text_audit(raw, column, use_failure=False):
    index = openpyxl.utils.column_index_from_string(column) - 1
    entries = [normalize(row[index]) for row in raw if not use_failure or normalize(row[19]) == "Yes"]
    usable = [value for value in entries if value and value != PLACEHOLDER]
    counts = Counter(usable)
    patterns = {
        "email": re.compile(r"[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}"),
        "phone": re.compile(r"(?<!\d)(?:\+?91[\s.-]?)?[6-9](?:[\s.-]?\d){9}(?!\d)"),
        "url": re.compile(r"https?://\S+", re.I),
        "longIdentifier": re.compile(r"(?<!\w)\d{12,}(?!\w)"),
    }
    return {"column": column, "population": "failure" if use_failure else "all",
            "eligible": len(entries), "usable": len(usable), "missing": sum(not value for value in entries),
            "placeholdersInEligible": entries.count(PLACEHOLDER), "uniqueTexts": len(counts),
            "repeatedTextGroups": sum(count > 1 for count in counts.values()),
            "recordsInRepeatedGroups": sum(count for count in counts.values() if count > 1),
            "repetitionBeyondFirst": len(usable) - len(counts), "largestRepeat": max(counts.values(), default=0),
            "potentialPiiCounts": {name: sum(bool(pattern.search(value)) for value in usable)
                                   for name, pattern in patterns.items()},
            "piiMethod": "Pattern scan only; names and contextual identifiers require review. Text export is owned by NLP pipeline."}


def prepare(source, destination):
    source = source.resolve(strict=True)
    destination.mkdir(parents=True, exist_ok=True)
    workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
    sheet = workbook["Sheet1"]
    rows = list(sheet.values)
    workbook.close()
    headers, raw = rows[0], rows[1:]
    if len(headers) != 32 or any(len(row) != 32 for row in raw):
        raise ValueError("Expected 32 columns in the supplied questionnaire")
    metrics = [{"id": key, "label": label, "question": normalize(headers[openpyxl.utils.column_index_from_string(column) - 1]),
                "column": column, "population": population, **({"dimension": dimension} if dimension else {})}
               for key, label, column, population, dimension in METRIC_SPECS]
    records = []
    for index, row in enumerate(raw, start=2):
        provider = normalize(row[4])
        if provider not in PROVIDERS:
            raise ValueError(f"Unexpected provider at row {index}: {provider!r}")
        record = {"id": f"r{index - 1:04d}", "provider": provider, "services": selections(row[5]),
                  "tenure": normalize(row[6]), "selectionReason": normalize(row[7]),
                  "support": boolean(row[12], "support", index), "failure": boolean(row[19], "failure", index),
                  "switching": boolean(row[26], "switching", index),
                  "failureTypes": [] if normalize(row[20]) == "None" else selections(row[20]),
                  "remedy": normalize(row[23]).replace("Explantion", "Explanation"), "trust": normalize(row[25]),
                  "switchReasons": selections(row[27]),
                  "ratings": {key: rating(row[openpyxl.utils.column_index_from_string(column) - 1], column, index)
                              for key, _, column, _, _ in METRIC_SPECS}}
        records.append(record)
    alpha = {"reliability": cronbach(records, ["reliability", "coverage", "speed", "promise"]),
             "responsiveness": cronbach(records, ["access", "response", "resolution"])}
    eligible_counts = {key: {"all": sum(r[key] for r in records),
                            **{provider: sum(r[key] and r["provider"] == provider for r in records) for provider in PROVIDERS}}
                       for key in ("support", "failure", "switching")}
    audit = {"version": VERSION, "worksheet": "Sheet1", "range": f"A1:AF{len(rows)}", "n": len(records), "columns": 32,
             "providerCounts": dict(Counter(r["provider"] for r in records)), "eligible": eligible_counts,
             "duplicates": {"ids": len(raw) - len({row[0] for row in raw}),
                            "completeRecords": len(raw) - len(set(raw))},
             "missingByColumn": {openpyxl.utils.get_column_letter(i + 1): sum(row[i] is None or normalize(row[i]) == "" for row in raw)
                                 for i in range(32)},
             "ratingRange": {"minimum": min(v for r in records for v in r["ratings"].values() if v is not None),
                             "maximum": max(v for r in records for v in r["ratings"].values() if v is not None), "invalid": 0},
             "ineligiblePopulated": {metric["id"]: sum(not eligible(r, metric["id"]) and r["ratings"][metric["id"]] is not None for r in records)
                                     for metric in metrics if metric["population"] != "all"},
             "ineligibleRecoveryCategorical": {key: sum(not r["failure"] and bool(r[key]) for r in records) for key in ("remedy", "trust")},
             "categories": {key: dict(Counter(value for r in records for value in r[key])) if key in ("services", "failureTypes", "switchReasons")
                            else dict(Counter(r[key] for r in records)) for key in ("services", "tenure", "selectionReason", "failureTypes", "remedy", "trust", "switchReasons")},
             "text": {"incident": text_audit(raw, "V", True), "best": text_audit(raw, "AE"), "improve": text_audit(raw, "AF")},
             "incidentPlaceholders": sum(normalize(row[21]) == PLACEHOLDER for row in raw),
             "alpha": alpha,
             "empathyReconciliation": {"slideClaim": {"Airtel": 3.22, "Jio": 3.25, "status": "does not reproduce from workbook"},
                                       "allRecords": {p: summary([r["ratings"]["empathy"] for r in records if r["provider"] == p]) for p in PROVIDERS},
                                       "eligibleSupport": {p: summary(values(records, "empathy", p)) for p in PROVIDERS}},
             "normalization": ["Unicode NFKC and whitespace trim/collapse, including non-breaking spaces", "Split semicolon selections; drop empty tokens; deduplicate within respondent", "Explantion → Explanation (display only)", "Blank failure type → empty list", "Call drop Billing retained intact as an ambiguous category"],
             "normalizationMap": {"Explantion": "Explanation", "Brand Trust\u00a0": "Brand Trust", "Customer Support\u00a0": "Customer Support", "Installation\u00a0": "Installation", "Compensation\u00a0": "Compensation", "Technical Correction\u00a0": "Technical Correction", "Increased\u00a0": "Increased", "Coverage\u00a0": "Coverage"},
             "ambiguousCategory": {"label": "Call drop Billing", "respondents": sum("Call drop Billing" in r["failureTypes"] for r in records)},
             "privacy": {"exportedIdentifiers": "Sequential anonymous keys only; original ID, name, email and language omitted", "namesPopulated": sum(bool(normalize(row[2])) for row in raw), "nonAnonymousEmails": sum(bool(normalize(row[1])) and normalize(row[1]).lower() != "anonymous" for row in raw)},
             "provenance": "Calculated from the supplied workbook; customer-response origin confirmed by the project owner.",
             "limitations": ["Recruitment, collection dates, population coverage and original questionnaire labels were not supplied.", "Scale direction is project-owner confirmed: 1 low, 5 high.", "The workbook measures perceptions without paired expectations; no classical SERVQUAL gaps can be calculated.", "Support N–R measures require support contact; recovery W/Y, remedy X and trust Z require failure.", "Potential switching reasons are hypothetical and populated for all records.", "Repeated text does not establish duplicate respondents.", "The raw workbook is retained unchanged outside the website export."]}
    segment_specs = [("all", {}), ("support", {"support": True}), ("failure", {"failure": True}),
                     ("switching", {"switching": True}), ("long-tenure", {"tenure": ["> 5 Years"]}),
                     ("prepaid-or-fiber", {"services": ["Prepaid", "Fiber Broadband"]}),
                     ("support-no-failure", {"support": True, "failure": False})]
    segments = []
    for segment_id, filters in segment_specs:
        selected = [r for r in records if matches(r, filters)]
        results = {key: comparison(values(selected, key, "Airtel"), values(selected, key, "Jio")) for key, *_ in METRIC_SPECS}
        segments.append({"id": segment_id, "filters": filters, "n": len(selected), "metrics": results, "holm": holm(results)})
    rng = np.random.default_rng(SEED)
    a, b = [1, 2, 2, 4, 5], [1, 3, 3, 4, 4, 5]
    indices_a = rng.integers(0, len(a), (40, len(a)))
    indices_b = rng.integers(0, len(b), (40, len(b)))
    differences = np.asarray(a)[indices_a].mean(axis=1) - np.asarray(b)[indices_b].mean(axis=1)
    fixtures = {"version": VERSION, "summaryConvention": "Sample variance/SD ddof=1; all tied modes; bins 1–5; proportions 0–1",
                "comparisonConvention": "Airtel minus Jio; asymptotic two-sided Mann–Whitney U with tie and continuity correction; p-values exist in fixtures for mathematical checks even if n<20 makes display inference unavailable",
                "holmConvention": "Five declared primary proxies; missing tests remain in family size; sort available p ascending and cumulative maximum of (5-index)*p capped at 1",
                "segments": segments, "alpha": alpha,
                "relationships": [relationship(records, *pair) for pair in [("reliability", "recommend"), ("response", "empathy"), ("recovery", "continue"), ("fairness", "recommend")]],
                "categorical": [categorical(records, field) for field in ("support", "failure", "switching")],
                "edgeCases": [{"id": "ties", "a": [1, 1, 3, 5], "b": [1, 2, 3, 5, 5], "expected": comparison([1, 1, 3, 5], [1, 2, 3, 5, 5])},
                              {"id": "constant", "a": [3] * 20, "b": [3] * 20, "expected": comparison([3] * 20, [3] * 20)},
                              {"id": "one-provider", "a": [1, 2, 3], "b": [], "expected": comparison([1, 2, 3], [])},
                              {"id": "empty", "a": [], "b": [], "expected": comparison([], [])},
                              {"id": "all-tied-modes", "values": [1, 2, 3, 4, 5], "expected": summary([1, 2, 3, 4, 5])}],
                "bootstrapFixture": {"seed": SEED, "a": a, "b": b, "indicesA": indices_a.tolist(), "indicesB": indices_b.tolist(),
                                     "interval": np.quantile(differences, [0.025, 0.975], method="linear").tolist(), "percentileMethod": "linear", "resamples": 40,
                                     "note": "Fixed indices isolate bootstrap/quantile checking from RNG implementation; production uses 2,000 resamples."}}
    output = {"survey.json": records, "audit.json": audit, "fixtures.json": fixtures, "metrics.json": metrics}
    for name, data in output.items():
        (destination / name).write_text(json.dumps(data, ensure_ascii=False, separators=(",", ":"), allow_nan=False), encoding="utf-8")
    manifest = {"version": VERSION, "source": {"filename": source.name, "sha256": hashlib.sha256(source.read_bytes()).hexdigest(), "sheet": "Sheet1", "range": audit["range"]},
                "transform": {"script": "prepare_survey.py", "sha256": hashlib.sha256(Path(__file__).read_bytes()).hexdigest()},
                "dependencies": {"python": platform.python_version(), **{name: importlib.metadata.version(name) for name in ("openpyxl", "numpy", "scipy")}},
                "resampling": {"seed": SEED, "bootstrapResamples": 2000, "spearmanPermutations": 4999, "confidenceLevel": 0.95, "bootstrapQuantile": "linear"},
                "assets": {name: hashlib.sha256((destination / name).read_bytes()).hexdigest() for name in output},
                "collectionDate": None, "samplingFrame": None, "provenance": audit["provenance"]}
    (destination / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    self_check(records, audit, fixtures)
    print(json.dumps({"records": len(records), "providers": audit["providerCounts"], "eligible": eligible_counts, "alpha": alpha, "output": str(destination)}, indent=2))


def self_check(records, audit, fixtures):
    """Small runnable regression check for the supplied source and transformations."""
    assert len(records) == 946 and audit["providerCounts"] == {"Jio": 513, "Airtel": 433}
    assert audit["eligible"]["support"] == {"all": 599, "Airtel": 271, "Jio": 328}
    assert audit["eligible"]["failure"] == {"all": 520, "Airtel": 240, "Jio": 280}
    assert audit["duplicates"] == {"ids": 0, "completeRecords": 0}
    assert audit["incidentPlaceholders"] == 426
    assert audit["ineligiblePopulated"]["response"] == 347 and audit["ineligiblePopulated"]["recovery"] == 426
    assert selections("Prepaid; Prepaid;Fiber Broadband; ") == ["Prepaid", "Fiber Broadband"]
    assert normalize("Coverage\u00a0") == "Coverage" and selections("Call drop Billing;") == ["Call drop Billing"]
    assert summary([])["mean"] is None and summary([2])["variance"] is None
    assert summary([1, 1, 2, 2, 3])["modes"] == [1, 2]
    assert comparison([3] * 20, [3] * 20)["p"] is None
    assert all(set(r) == {"id", "provider", "services", "tenure", "selectionReason", "support", "failure", "switching", "failureTypes", "remedy", "trust", "switchReasons", "ratings"} for r in records)
    assert not any("\u00a0" in json.dumps(r, ensure_ascii=False) for r in records)
    assert len({r["id"] for r in records}) == len(records)
    assert abs(audit["alpha"]["reliability"]["alpha"] - 0.474) < 0.001
    assert abs(audit["alpha"]["responsiveness"]["alpha"] - 0.401) < 0.001
    assert fixtures["segments"][0]["metrics"]["empathy"]["Airtel"]["n"] == 271


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()
    prepare(args.input, args.output)
